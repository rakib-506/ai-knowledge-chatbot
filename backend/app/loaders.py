"""Turn different file formats (and web pages) into plain text."""
import io
import re
from pathlib import Path

import pymupdf as fitz           # PyMuPDF, for PDF
import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument


def clean_text(text: str) -> str:
    text = text.replace("\u200b", "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)     # keep paragraph breaks
    return text.strip()


def load_pdf(data: bytes) -> str:
    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text("text") for page in doc)


def load_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                       # also read tables
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg", "form"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else ""
    body = soup.get_text(separator="\n")
    return f"{title}\n\n{body}" if title else body


def load_file(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = load_pdf(data)
    elif ext == ".docx":
        text = load_docx(data)
    elif ext in (".html", ".htm"):
        text = html_to_text(data.decode("utf-8", errors="replace"))
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"File type '{ext}' is not supported.")
    return clean_text(text)


def load_url(url: str) -> str:
    resp = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0 (KnowledgeChatbot)"})
    resp.raise_for_status()
    return clean_text(html_to_text(resp.text))
